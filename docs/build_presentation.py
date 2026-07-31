"""Génère le support de présentation My Content (.pptx, ~19 slides).

Reproductible et versionné (démarche « industrialisable »).
Sortie : docs/My_Content_presentation.pptx
Export PDF : ouvrir dans PowerPoint -> Fichier -> Exporter -> PDF.

Usage : python docs/build_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --- Palette --------------------------------------------------------------
PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)   # bleu nuit
ACCENT = RGBColor(0x2E, 0x86, 0xC1)    # bleu
ACCENT2 = RGBColor(0xE6, 0x7E, 0x22)   # orange
GREEN = RGBColor(0x27, 0xAE, 0x60)
RED = RGBColor(0xC0, 0x39, 0x2B)
LIGHT = RGBColor(0xF4, 0xF6, 0xF7)
DARK = RGBColor(0x2C, 0x3E, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x7F, 0x8C, 0x8D)

SW, SH = Inches(13.333), Inches(7.5)


def _set_text(tf, size, color, bold=False, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.runs[0] if p.runs else p.add_run()
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold


def title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(1.15))
    bar.fill.solid(); bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()
    tf = bar.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.text = title
    _set_text(tf, 26, WHITE, bold=True)
    # bande d'accent
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.15), SW, Inches(0.07))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT2
    strip.line.fill.background()
    if subtitle:
        tb = slide.shapes.add_textbox(Inches(0.4), Inches(1.25), Inches(12.5), Inches(0.5))
        tb.text_frame.text = subtitle
        _set_text(tb.text_frame, 15, GREY, bold=False)


def bullets(slide, items, left=0.6, top=1.9, width=12.1, height=5.0, size=17):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        text, level = item if isinstance(item, tuple) else (item, 0)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = level
        r = p.add_run(); r.text = ("• " if level == 0 else "– ") + text
        r.font.size = Pt(size - 2 * level)
        r.font.color.rgb = DARK if level == 0 else GREY
        p.space_after = Pt(8)
    return tb


def box(slide, x, y, w, h, text, fill=ACCENT, font=12, tcolor=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.color.rgb = PRIMARY; sp.line.width = Pt(1)
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.text = text
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size = Pt(font); r.font.color.rgb = tcolor; r.font.bold = True
    return sp


def arrow(slide, x, y, w, h, shape=MSO_SHAPE.RIGHT_ARROW, fill=GREY):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    sp.line.fill.background()
    return sp


def label(slide, x, y, w, text, size=10, color=GREY):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.35))
    tb.text_frame.text = text
    _set_text(tb.text_frame, size, color, align=PP_ALIGN.CENTER)


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def pros_cons(slide, pros, cons, top=2.1):
    for title, items, color, x in [("Avantages", pros, GREEN, 0.6),
                                    ("Inconvénients", cons, RED, 6.9)]:
        head = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top),
                                      Inches(5.8), Inches(0.55), )
        head.fill.solid(); head.fill.fore_color.rgb = color; head.line.fill.background()
        head.text_frame.text = title
        _set_text(head.text_frame, 16, WHITE, bold=True, align=PP_ALIGN.CENTER)
        bullets(slide, items, left=x + 0.1, top=top + 0.75, width=5.6, height=3.6, size=14)


def table(slide, rows, top=2.0, left=0.5, width=12.3, height=4.5, header_fill=PRIMARY):
    n, m = len(rows), len(rows[0])
    gtbl = slide.shapes.add_table(n, m, Inches(left), Inches(top),
                                  Inches(width), Inches(height)).table
    for j, cell in enumerate(rows[0]):
        c = gtbl.cell(0, j); c.text = cell
        c.fill.solid(); c.fill.fore_color.rgb = header_fill
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = WHITE
    for i in range(1, n):
        for j in range(m):
            c = gtbl.cell(i, j); c.text = rows[i][j]
            c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(11); r.font.color.rgb = DARK
    return gtbl


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # 1 — Titre
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = PRIMARY; bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(1), Inches(2.4), Inches(11.3), Inches(1.5))
    t.text_frame.text = "My Content"
    _set_text(t.text_frame, 54, WHITE, bold=True)
    st = s.shapes.add_textbox(Inches(1), Inches(3.7), Inches(11.3), Inches(1))
    st.text_frame.text = "MVP — Système de recommandation d'articles"
    _set_text(st.text_frame, 26, RGBColor(0xAE, 0xD6, 0xF1))
    strip = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(4.6), Inches(3), Inches(0.08))
    strip.fill.solid(); strip.fill.fore_color.rgb = ACCENT2; strip.line.fill.background()
    ft = s.shapes.add_textbox(Inches(1), Inches(6.4), Inches(11), Inches(0.5))
    ft.text_frame.text = "Présentation à Samia (CEO) — support technique & architecture"
    _set_text(ft.text_frame, 14, RGBColor(0xAE, 0xD6, 0xF1))

    # 2 — Contexte & objectif
    s = blank(prs); title_bar(s, "Contexte & objectif")
    bullets(s, [
        "My Content : start-up qui veut encourager la lecture par la recommandation de contenus.",
        "Pas encore de données propriétaires → jeu de données public (interactions Globo.com).",
        ("Fonctionnalité MVP la plus critique :", 0),
        ("« En tant qu'utilisateur, je reçois une sélection de 5 articles. »", 1),
        "Contrainte structurante : anticiper l'ajout de nouveaux utilisateurs et de nouveaux articles.",
        "Exigence : solution serverless, industrialisable, déployable de bout en bout.",
    ])

    # 3 — Description fonctionnelle
    s = blank(prs); title_bar(s, "Description fonctionnelle de l'application")
    steps = [("Utilisateur\n(id)", ACCENT), ("Application\n(Streamlit)", ACCENT),
             ("Service de reco\n(Azure Function)", ACCENT2), ("5 articles\nrecommandés", GREEN)]
    x = 0.7
    for i, (txt, col) in enumerate(steps):
        box(s, x, 2.8, 2.6, 1.5, txt, fill=col, font=15)
        x += 2.6
        if i < len(steps) - 1:
            arrow(s, x - 0.05, 3.35, 0.55, 0.4)
            x += 0.55
    bullets(s, [
        "L'app liste des identifiants utilisateurs et déclenche l'appel du service.",
        "Le service serverless reçoit un user_id et renvoie le top-5 des articles.",
        "Objectif : démontrer la fonctionnalité à Samia et à de futurs utilisateurs.",
    ], top=4.7)

    # 4 — Les données
    s = blank(prs); title_bar(s, "Les données", "Globo.com — interactions utilisateurs ↔ articles")
    bullets(s, [
        ("Clics / sessions : user_id, session, article cliqué, horodatage.", 0),
        ("Articles : catégorie, date de création, éditeur, nombre de mots.", 0),
        ("Embeddings : vecteur de 250 dimensions par article (contenu).", 0),
        ("Nature : feedback implicite (clics), pas de notes explicites.", 0),
        ("Enjeux : forte creusité (users/articles peu actifs), volumétrie des embeddings.", 0),
    ])

    # 5 — Approches de modélisation (vue d'ensemble)
    s = blank(prs); title_bar(s, "Approches de recommandation analysées")
    for x, name, col, desc in [
        (0.7, "Content-based", ACCENT, "Similarité de contenu\n(embeddings d'articles)"),
        (4.9, "Collaborative\nfiltering (ALS)", ACCENT2, "Goûts d'utilisateurs\nsimilaires"),
        (9.1, "Hybride", GREEN, "Combinaison des deux\n+ fallback popularité"),
    ]:
        box(s, x, 2.4, 3.5, 1.3, name, fill=col, font=17)
        b = s.shapes.add_textbox(Inches(x), Inches(3.9), Inches(3.5), Inches(1.2))
        b.text_frame.text = desc; b.text_frame.word_wrap = True
        _set_text(b.text_frame, 13, DARK, align=PP_ALIGN.CENTER)
    bullets(s, ["Les trois stratégies sont exposées derrière un même service ; l'hybride est la stratégie par défaut."],
            top=5.4)

    # 6 — Content-based
    s = blank(prs); title_bar(s, "Modèle 1 — Filtrage basé sur le contenu")
    bullets(s, ["Principe : profil utilisateur = moyenne des embeddings des articles lus, "
                "puis similarité cosinus vers tout le catalogue."], top=1.7, height=0.9)
    pros_cons(s,
              ["Fonctionne dès le 1ᵉʳ clic (peu de cold start utilisateur)",
               "Intègre un nouvel article immédiatement (il a un embedding)",
               "Explicable : « proche de ce que vous avez lu »"],
              ["N'exploite pas l'intelligence collective",
               "Risque de « bulle » (peu de diversité/sérendipité)",
               "Dépend de la qualité des embeddings"])

    # 7 — Collaborative
    s = blank(prs); title_bar(s, "Modèle 2 — Filtrage collaboratif (ALS)")
    bullets(s, ["Principe : factorisation de la matrice utilisateur × article "
                "(feedback implicite) par ALS → facteurs latents."], top=1.7, height=0.9)
    pros_cons(s,
              ["Capte des goûts latents non visibles dans le contenu",
               "Effet de découverte / sérendipité",
               "Performant sur les utilisateurs actifs"],
              ["Cold start fort (nouvel utilisateur ET nouvel article)",
               "Ré-entraînement nécessaire (ALS non incrémental)",
               "Sensible à la creusité des données"])

    # 8 — Hybride
    s = blank(prs); title_bar(s, "Modèle 3 — Hybride + cold start")
    bullets(s, ["Principe : combinaison des scores content + collaboratif (normalisés) ; "
                "repli sur les articles populaires si historique insuffisant."], top=1.7, height=0.9)
    pros_cons(s,
              ["Combine les forces des deux approches",
               "Dégradation propre : popularité en dernier recours",
               "Couvre nouveaux utilisateurs et nouveaux articles"],
              ["Deux modèles à maintenir et à synchroniser",
               "Pondération à régler (paramètre alpha)",
               "Un peu plus de complexité opérationnelle"])

    # 9 — Comparaison
    s = blank(prs); title_bar(s, "Comparaison des modèles")
    table(s, [
        ["Critère", "Content-based", "Collaboratif (ALS)", "Hybride"],
        ["Nouvel utilisateur", "Bon (dès 1 clic)", "Faible", "Bon"],
        ["Nouvel article", "Excellent", "Faible", "Bon"],
        ["Sérendipité", "Faible", "Bonne", "Bonne"],
        ["Coût d'entraînement", "Faible", "Élevé (batch)", "Moyen"],
        ["Explicabilité", "Élevée", "Faible", "Moyenne"],
        ["Choix MVP", "—", "—", "✔ retenu"],
    ], height=4.3)
    bullets(s, ["Évaluation par leave-last-out (HitRate@5) dans le notebook pour départager sur nos données."],
            top=6.4, height=0.6, size=13)

    # 10 — Système retenu
    s = blank(prs); title_bar(s, "Système de recommandation retenu")
    bullets(s, [
        "Stratégie hybride par défaut, exposée via le service serverless.",
        ("Chaîne de décision pour un user_id :", 0),
        ("1. Historique exploitable → score hybride (content + collaboratif)", 1),
        ("2. Historique faible → content-based seul", 1),
        ("3. Aucun historique → articles les plus populaires (cold start)", 1),
        "Inférence légère : ne dépend que de numpy (artefacts pré-calculés hors-ligne).",
        "Renvoie le top-5 (paramétrable).",
    ])

    # 11 — Astuce ACP
    s = blank(prs); title_bar(s, "Optimisation production — réduction de dimension (ACP)")
    box(s, 1.3, 3.0, 3.6, 1.6, "Embeddings\n250 dimensions", fill=ACCENT, font=16)
    arrow(s, 5.1, 3.6, 1.2, 0.4)
    box(s, 6.6, 3.0, 3.0, 1.6, "ACP", fill=ACCENT2, font=18)
    arrow(s, 9.8, 3.6, 1.2, 0.4)
    box(s, 11.0, 3.0, 1.9, 1.6, "~50 dim", fill=GREEN, font=16)
    bullets(s, [
        "Fichier d'embeddings volumineux vs limites du free tier Azure.",
        "Réduction par ACP hors-ligne → artefact plus léger, latence réduite.",
        "Compromis maîtrisé entre taille/temps et qualité des recommandations.",
    ], top=4.9)

    # 12 — Architecture retenue (deux solutions indépendantes)
    s = blank(prs); title_bar(s, "Architecture retenue — deux solutions indépendantes",
                              "Même cœur de reco et mêmes artefacts, deux déploiements autonomes")
    # ① Azure (serverless / industrialisable)
    label(s, 0.6, 1.65, 5.0, "① Solution Azure — serverless", size=13, color=ACCENT2)
    box(s, 0.6, 2.05, 2.5, 1.0, "App Streamlit", fill=ACCENT, font=13)
    arrow(s, 3.15, 2.4, 0.9, 0.35); label(s, 3.0, 2.05, 1.3, "user_id")
    box(s, 4.15, 2.05, 3.0, 1.0, "Azure Function\n(numpy)", fill=ACCENT2, font=13)
    arrow(s, 7.2, 2.4, 0.9, 0.35)
    box(s, 8.2, 2.05, 3.0, 1.0, "Azure Blob\n(modèles)", fill=PRIMARY, font=13)
    # ② Hugging Face (démo publique auto-suffisante)
    label(s, 0.6, 3.35, 6.0, "② Solution Hugging Face — démo publique", size=13, color=ACCENT)
    box(s, 0.6, 3.75, 3.7, 1.0, "HF Space (Streamlit\n+ Recommender)", fill=ACCENT, font=13)
    arrow(s, 4.35, 4.1, 0.9, 0.35); label(s, 4.2, 3.75, 1.3, "charge")
    box(s, 5.35, 3.75, 3.0, 1.0, "HF Hub\n(modèles)", fill=PRIMARY, font=13)
    # Pipeline commun
    box(s, 1.5, 5.5, 10.3, 1.0,
        "Pipeline hors-ligne COMMUN (prepare_model.py) : données brutes → ACP + ALS "
        "→ artefacts publiés vers Blob ET HF Hub", fill=GREY, font=13)

    # 13 — Justification du choix d'archi
    s = blank(prs); title_bar(s, "Pourquoi l'Architecture 2 pour le MVP ?")
    table(s, [
        ["Critère", "Archi 1 — API dédiée", "Archi 2 — Blob (retenue)"],
        ["Pièces mobiles", "API + Function", "Function seule"],
        ["Coût (free tier)", "Plus élevé", "Quasi nul"],
        ["Mise en place", "Plus longue", "Rapide"],
        ["Séparation des responsabilités", "Meilleure", "Suffisante (MVP)"],
        ["Montée en charge", "Meilleure (cible)", "Limitée"],
    ], height=3.9)
    bullets(s, ["MVP : simplicité et coût priment → Archi 2. L'Archi 1 est visée pour l'architecture cible."],
            top=6.3, height=0.6, size=13)

    # 14 — Démo application
    s = blank(prs); title_bar(s, "Démonstration — les applications")
    bullets(s, [
        "Interface Streamlit : sélection d'un identifiant utilisateur → 5 articles.",
        ("Deux vitrines, selon la solution :", 0),
        ("Azure : app locale qui appelle le service serverless (Azure Function).", 1),
        ("Hugging Face : Space public auto-suffisant (reco calculée sur place).", 1),
        "Paramétrable : stratégie de reco, nombre d'articles.",
    ], width=7.9)
    box(s, 8.9, 2.2, 3.8, 3.6,
        "[ Capture d'écran\ndu Space HF\nà insérer ]", fill=LIGHT, font=14, tcolor=GREY,
        shape=MSO_SHAPE.RECTANGLE)

    # 15 — Industrialisation
    s = blank(prs); title_bar(s, "Industrialisation & déploiement bout-en-bout")
    bullets(s, [
        "Code versionné : Git local + push GitHub.",
        "Séparation claire : préparation hors-ligne / service d'inférence / application.",
        ("Déploiement reproductible, deux cibles depuis les mêmes artefacts :", 0),
        ("prepare_model.py → artefacts (source commune)", 1),
        ("Azure : func azure functionapp publish (+ artefacts vers Blob)", 1),
        ("Hugging Face : push du Space + upload_artifacts.py vers le HF Hub", 1),
        "Support de présentation lui-même généré par script (reproductible).",
    ])

    # 16 — Architecture cible (schéma)
    s = blank(prs); title_bar(s, "Architecture cible — nouveaux utilisateurs & articles",
                              "Intégration en continu, sans tout recalculer")
    box(s, 0.5, 2.1, 2.2, 1.0, "Clics\n(temps réel)", fill=ACCENT, font=12)
    arrow(s, 2.75, 2.45, 0.7, 0.35)
    box(s, 3.5, 2.1, 2.4, 1.0, "Event Hub /\nFunction ingestion", fill=ACCENT2, font=12)
    arrow(s, 5.95, 2.45, 0.7, 0.35)
    box(s, 6.7, 2.1, 2.4, 1.0, "Cosmos DB\n(profils users)", fill=PRIMARY, font=12)
    box(s, 0.5, 3.7, 3.3, 1.0, "Nouvel article →\nembedding + ACP", fill=ACCENT, font=12)
    box(s, 4.2, 3.7, 3.6, 1.0, "Ré-entraînement ALS\nplanifié (Azure ML / timer)", fill=ACCENT2, font=12)
    box(s, 9.4, 2.9, 3.3, 1.0, "Catalogue + modèles\n(Blob, versionnés)", fill=PRIMARY, font=12)
    arrow(s, 7.9, 4.1, 1.5, 0.35); arrow(s, 3.8, 4.1, 0.4, 0.35)
    box(s, 9.4, 4.3, 3.3, 1.0, "API de reco dédiée\n(Architecture 1)", fill=GREEN, font=12)
    arrow(s, 10.9, 3.95, 0.4, 0.35, shape=MSO_SHAPE.DOWN_ARROW)
    box(s, 5.5, 5.9, 3.3, 0.9, "Applications clientes", fill=ACCENT, font=13)
    arrow(s, 9.7, 5.35, 1.2, 0.4, shape=MSO_SHAPE.LEFT_ARROW)
    label(s, 0.5, 6.2, 4.5, "Suivi : Application Insights + métriques métier (CTR)", size=11, color=GREY)

    # 17 — Cold start dans la cible
    s = blank(prs); title_bar(s, "Cible — comment sont gérés les nouveaux venus")
    bullets(s, [
        ("Nouvel article :", 0),
        ("calcul de son embedding (+ ACP) → disponible immédiatement via le content-based, sans ré-entraînement.", 1),
        ("Nouvel utilisateur :", 0),
        ("popularité au tout début, puis content-based dès le 1ᵉʳ clic ; profil enrichi en continu via l'ingestion.", 1),
        ("Modèle collaboratif :", 0),
        ("ré-entraînement planifié (batch) ; entre deux, repli propre sur content-based / popularité.", 1),
        ("Montée en charge :", 0),
        ("passage à l'API dédiée (Archi 1) : mise à l'échelle, cache, A/B testing, versionnage des modèles.", 1),
    ], size=16)

    # 18 — Limites & suites
    s = blank(prs); title_bar(s, "Limites connues & prochaines étapes")
    bullets(s, [
        "Évaluation à approfondir : MAP@k, couverture, diversité (au-delà du HitRate@5).",
        "Prise en compte de la fraîcheur / récence des articles.",
        "Sécurité : clé de fonction (MVP) → authentification robuste à terme.",
        "Passage progressif à l'ingestion événementielle et au ré-entraînement automatisé.",
        "Collecte de nos propres données d'usage pour affiner les modèles.",
    ])

    # 19 — Conclusion
    s = blank(prs)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid(); bg.fill.fore_color.rgb = PRIMARY; bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11.3), Inches(1.5))
    t.text_frame.text = "MVP fonctionnel, serverless et industrialisable"
    _set_text(t.text_frame, 34, WHITE, bold=True)
    st = s.shapes.add_textbox(Inches(1), Inches(4.1), Inches(11.3), Inches(1))
    st.text_frame.text = "Prêt à démontrer les 5 recommandations et à évoluer vers l'architecture cible."
    _set_text(st.text_frame, 18, RGBColor(0xAE, 0xD6, 0xF1))

    out = Path(__file__).resolve().parent / "My_Content_presentation.pptx"
    prs.save(out)
    print(f"Présentation générée : {out}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build()
